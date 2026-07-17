# ─────────────────────────────────────────────────────────────────────────────
#  GENERADOR DE EXCEL
# ─────────────────────────────────────────────────────────────────────────────

import io
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generar_excel(hojas: dict) -> bytes:
    """
    Recibe dict { nombre_hoja: [movimientos] }
    Devuelve el archivo Excel como bytes (para descarga).

    Estructura de cada hoja:
      Fila 1: encabezados
      Fila 2: SALDO INICIAL (derivado del primer movimiento = el más antiguo,
              ya que la app ordena los movimientos cronológicamente)
      Fila 3+: movimientos, con dos columnas de control calculadas:
              - "Saldo calculado": saldo que resulta de arrastrar el saldo inicial
                aplicando créditos y débitos (fórmula de Excel, se recalcula sola)
              - "Diferencia": saldo calculado − saldo informado (0 si todo cuadra)
      La columna "Cuenta" solo se agrega si la hoja tiene ese dato (BBVA).
    """
    wb = openpyxl.Workbook()
    if wb.sheetnames == ["Sheet"]:
        del wb["Sheet"]

    # openpyxl no acepta en el título de hoja los caracteres  \ / * ? : [ ]
    # ni títulos vacíos o > 31 chars, y no permite duplicados. Se sanea el
    # nombre que puso el usuario para que se aplique igual (sin romper).
    _usados = set()

    def _titulo_hoja(nombre: str) -> str:
        limpio = re.sub(r"[\\/*?:\[\]]", "-", (nombre or "").strip())[:31].strip()
        if not limpio:
            limpio = "Hoja"
        base = limpio
        n = 2
        while limpio.lower() in _usados:
            sufijo = f" ({n})"
            limpio = base[:31 - len(sufijo)] + sufijo
            n += 1
        _usados.add(limpio.lower())
        return limpio

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    data_font   = Font(name="Arial", size=9)
    bold_font   = Font(name="Arial", size=9, bold=True)
    alt_fill    = PatternFill("solid", fgColor="EBF3FB")
    white_fill  = PatternFill("solid", fgColor="FFFFFF")
    inicial_fill = PatternFill("solid", fgColor="FFF2CC")  # fila SALDO INICIAL
    thin        = Side(style="thin", color="CCCCCC")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    num_fmt     = "#,##0.00"

    # Columnas fijas (A..I). La columna "Cuenta" se agrega al final si corresponde.
    # Índices (1-based):
    #   1 Nro de Mes | 2 Fecha | 3 Descripcion | 4 (vacía) | 5 Debito | 6 Credito
    #   7 Saldo | 8 Saldo calculado | 9 Diferencia | (10 Cuenta, opcional)
    headers_base = ["Nro de Mes", "Fecha", "Descripcion", "",
                    "Debito", "Credito", "Saldo", "Saldo calculado", "Diferencia"]
    widths_base  = [12, 14, 60, 18, 16, 16, 20, 18, 16]
    COL_CUENTA = 10

    for sheet_name, txs in hojas.items():
        ws = wb.create_sheet(title=_titulo_hoja(sheet_name))

        tiene_cuenta = any(tx.get("cuenta") for tx in txs)
        headers = list(headers_base)
        widths  = list(widths_base)
        if tiene_cuenta:
            headers.append("Cuenta")
            widths.append(18)

        # ── Encabezados (fila 1) ──
        for col, (hdr, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=1, column=col, value=hdr)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = border
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.row_dimensions[1].height = 20

        # ── Fila 2: SALDO INICIAL ──
        # Preferimos el saldo inicial que informa el banco en el extracto más
        # antiguo (los movimientos vienen ordenados, así que txs[0] es el más
        # viejo y trae el saldo inicial de su informe). Si el banco no lo rotula
        # (ej. Galicia), lo derivamos del primer movimiento:
        #   saldo_inicial = saldo_1er_mov − credito_1er_mov + debito_1er_mov
        saldo_inicial = None
        if txs:
            reportado = txs[0].get("saldo_inicial")
            if reportado is not None:
                saldo_inicial = reportado
            elif txs[0]["saldo"] is not None:
                primero = txs[0]
                saldo_inicial = round(
                    primero["saldo"]
                    - (primero["credito"] or 0)
                    + (primero["debito"] or 0),
                    2,
                )
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=2, column=col)
            cell.fill   = inicial_fill
            cell.border = border
            if col == 3:
                cell.value     = "SALDO INICIAL"
                cell.font      = bold_font
                cell.alignment = Alignment(horizontal="left")
            elif col == 7:
                cell.value        = saldo_inicial
                cell.font         = bold_font
                cell.alignment    = Alignment(horizontal="right")
                if saldo_inicial is not None:
                    cell.number_format = num_fmt
            else:
                cell.font = data_font

        # ── Movimientos (fila 3 en adelante) ──
        for i, tx in enumerate(txs):
            r = i + 3  # los datos arrancan en la fila 3
            fill = alt_fill if r % 2 == 0 else white_fill

            valores = [
                tx["mes"], tx["fecha"], tx["descripcion"], None,
                tx["debito"], tx["credito"], tx["saldo"],
                # Saldo calculado: arrastra el saldo inicial (G2) fila a fila
                (f"=G2+F{r}-E{r}" if r == 3 else f"=H{r-1}+F{r}-E{r}"),
                # Diferencia: saldo calculado − saldo informado
                f"=H{r}-G{r}",
            ]
            if tiene_cuenta:
                valores.append(tx.get("cuenta"))

            for col, val in enumerate(valores, 1):
                cell = ws.cell(row=r, column=col, value=val)
                cell.font   = data_font
                cell.fill   = fill
                cell.border = border
                if col <= 2:
                    cell.alignment = Alignment(horizontal="center")
                elif col == 3:
                    cell.alignment = Alignment(horizontal="left")
                elif col == 4:
                    pass  # columna vacía
                elif col == COL_CUENTA:
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="right")
                    # Formato numérico para Debito/Credito/Saldo/Saldo calculado/Diferencia
                    if val is not None:
                        cell.number_format = num_fmt

        ws.freeze_panes = "A3"  # deja fijos encabezado + SALDO INICIAL

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
