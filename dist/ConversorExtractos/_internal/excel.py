# ─────────────────────────────────────────────────────────────────────────────
#  GENERADOR DE EXCEL
# ─────────────────────────────────────────────────────────────────────────────

import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def generar_excel(hojas: dict) -> bytes:
    """
    Recibe dict { nombre_hoja: [movimientos] }
    Devuelve el archivo Excel como bytes (para descarga).
    """
    wb = openpyxl.Workbook()
    if wb.sheetnames == ["Sheet"]:
        del wb["Sheet"]

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    data_font   = Font(name="Arial", size=9)
    alt_fill    = PatternFill("solid", fgColor="EBF3FB")
    white_fill  = PatternFill("solid", fgColor="FFFFFF")
    thin        = Side(style="thin", color="CCCCCC")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    num_fmt     = "#,##0.00"
    headers     = ["Nro de Mes", "Fecha", "Descripcion", "" , "Debito", "Credito", "Saldo"]
    col_widths  = [12, 14, 60, 18, 18, 20]

    for sheet_name, txs in hojas.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        for col, (hdr, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col, value=hdr)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = border
            ws.column_dimensions[cell.column_letter].width = w
        ws.row_dimensions[1].height = 20
        for row_idx, tx in enumerate(txs, 2):
            fill = alt_fill if row_idx % 2 == 0 else white_fill
            valores = [
                tx["mes"], tx["fecha"], tx["descripcion"], None,
                tx["debito"], tx["credito"], tx["saldo"],
            ]
            for col, val in enumerate(valores, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.font   = data_font
                cell.fill   = fill
                cell.border = border
                if col <= 2:
                    cell.alignment = Alignment(horizontal="center")
                elif col == 3:
                    cell.alignment = Alignment(horizontal="left")
                elif col == 4:
                    pass  # columna vacía, no hacer nada
                else:
                    cell.alignment = Alignment(horizontal="right")
                    if val is not None:
                        cell.number_format = num_fmt
        ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
